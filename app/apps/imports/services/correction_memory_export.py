from __future__ import annotations

import json
import tempfile
import zipfile
from itertools import islice

from django.http import FileResponse
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font


DEFAULT_ROWS_PER_FILE = 25_000
EXCEL_CELL_LIMIT = 32_767

EXPORT_COLUMNS = (
    ('Memory ID', 'memory_id'),
    ('Workflow', 'workflow'),
    ('Client', 'client'),
    ('SKU / record key', 'record_key'),
    ('Active', 'is_active'),
    ('Source fingerprint', 'source_fingerprint'),
    ('Solution fingerprint', 'proposal_fingerprint'),
    ('Origin file ID', 'origin_file_id'),
    ('Origin filename', 'origin_filename'),
    ('Approved by', 'approved_by'),
    ('Approved at', 'approved_at'),
    ('Last reused at', 'last_used_at'),
    ('Reuse count', 'use_count'),
    ('Approval reason', 'approval_note'),
    ('Approved action', 'approved_action'),
    ('Approved name', 'approved_name'),
    ('Approved description', 'approved_description'),
    ('Approved length (m)', 'approved_length_m'),
    ('Approved width (m)', 'approved_width_m'),
    ('Approved height (m)', 'approved_height_m'),
    ('Approved weight (kg)', 'approved_weight_kg'),
    ('Approved cubic (m3)', 'approved_cubic_m3'),
    ('Approved C/P', 'approved_freight_type'),
    ('Source name', 'source_name'),
    ('Source description', 'source_description'),
    ('Source category', 'source_category'),
    ('Source length (mm)', 'source_length_mm'),
    ('Source width (mm)', 'source_width_mm'),
    ('Source height (mm)', 'source_height_mm'),
    ('Source weight (kg)', 'source_weight_kg'),
    ('Source cubic (m3)', 'source_cubic_m3'),
    ('Source quantity', 'source_quantity'),
    ('Source pallet', 'source_pallet'),
    ('Source status', 'source_status'),
    ('Source comment', 'source_comment'),
    ('Field decisions JSON', 'field_decisions_json'),
    ('Source JSON', 'source_json'),
    ('Approved JSON', 'approved_json'),
)


def _safe_cell(value):
    if value is None:
        return ''
    if isinstance(value, (dict, list, tuple)):
        value = json.dumps(value, ensure_ascii=False, sort_keys=True, default=str)
    elif hasattr(value, 'isoformat'):
        value = value.isoformat()
    elif not isinstance(value, (str, int, float, bool)):
        value = str(value)
    if isinstance(value, str) and len(value) > EXCEL_CELL_LIMIT:
        return value[:EXCEL_CELL_LIMIT - 18] + '… [TRUNCATED]'
    return value


def _memory_values(memory):
    source_data = dict(memory.source_data or {})
    source = dict(source_data.get('source') or {})
    approved_data = dict(memory.approved_data or {})
    approved = dict(approved_data.get('approved_values') or {})
    values = {
        'memory_id': memory.pk,
        'workflow': memory.workflow_key,
        'client': memory.client.code,
        'record_key': memory.record_key,
        'is_active': memory.is_active,
        'source_fingerprint': memory.source_fingerprint,
        'proposal_fingerprint': memory.proposal_fingerprint,
        'origin_file_id': memory.origin_external_file_id,
        'origin_filename': (
            memory.origin_external_file.original_filename
            if memory.origin_external_file else ''
        ),
        'approved_by': (
            memory.approved_by.get_username() if memory.approved_by else ''
        ),
        'approved_at': memory.approved_at,
        'last_used_at': memory.last_used_at,
        'use_count': memory.use_count,
        'approval_note': memory.approval_note,
        'approved_action': approved_data.get('row_action'),
        'approved_name': approved.get('name'),
        'approved_description': approved.get('description'),
        'approved_length_m': approved.get('length_m'),
        'approved_width_m': approved.get('width_m'),
        'approved_height_m': approved.get('height_m'),
        'approved_weight_kg': approved.get('weight_kg'),
        'approved_cubic_m3': approved.get('cubic_m3'),
        'approved_freight_type': approved.get('freight_type'),
        'source_name': source.get('name'),
        'source_description': source.get('description'),
        'source_category': source.get('category'),
        'source_length_mm': source.get('length_mm'),
        'source_width_mm': source.get('width_mm'),
        'source_height_mm': source.get('height_mm'),
        'source_weight_kg': source.get('weight_kg'),
        'source_cubic_m3': source.get('cubic_m3'),
        'source_quantity': source.get('quantity'),
        'source_pallet': source.get('pallet'),
        'source_status': source.get('source_status'),
        'source_comment': source.get('comment'),
        'field_decisions_json': approved_data.get('field_decisions') or {},
        'source_json': source_data,
        'approved_json': approved_data,
    }
    return [_safe_cell(values[key]) for _label, key in EXPORT_COLUMNS]


def _header_row(worksheet, labels):
    cells = []
    for label in labels:
        cell = WriteOnlyCell(worksheet, value=label)
        cell.font = Font(bold=True)
        cells.append(cell)
    worksheet.append(cells)


def _build_workbook(memories, *, total_count, part_number, part_count, scope_label):
    workbook = Workbook(write_only=True)
    summary = workbook.create_sheet('Summary')
    _header_row(summary, ('Field', 'Value'))
    summary.append(('Export', 'Reconciliation Decision Memory'))
    summary.append(('Scope', scope_label))
    summary.append(('Generated at', timezone.localtime().isoformat()))
    summary.append(('Total records', total_count))
    summary.append(('This file part', f'{part_number} of {part_count}'))
    summary.append(('Safety', 'Report only; this export cannot update operational Products.'))

    sheet = workbook.create_sheet('Memories')
    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = f'A1:AL1'
    _header_row(sheet, [label for label, _key in EXPORT_COLUMNS])
    for memory in memories:
        sheet.append(_memory_values(memory))
    return workbook


def correction_memory_export_response(
    queryset,
    *,
    filename_prefix='reconciliation_memory',
    scope_label='All visible memories',
    rows_per_file=DEFAULT_ROWS_PER_FILE,
):
    """Create a disk-spooled Excel export; split large exports into a ZIP."""
    rows_per_file = max(1, int(rows_per_file))
    queryset = queryset.select_related(
        'client', 'approved_by', 'origin_external_file'
    ).order_by('pk')
    total_count = queryset.count()
    part_count = max(1, (total_count + rows_per_file - 1) // rows_per_file)
    stamp = timezone.localtime().strftime('%m%d.%H%M')
    iterator = queryset.iterator(chunk_size=1000)

    if part_count == 1:
        output = tempfile.SpooledTemporaryFile(max_size=8 * 1024 * 1024, mode='w+b')
        workbook = _build_workbook(
            iterator,
            total_count=total_count,
            part_number=1,
            part_count=1,
            scope_label=scope_label,
        )
        workbook.save(output)
        output.seek(0)
        return FileResponse(
            output,
            as_attachment=True,
            filename=f'{filename_prefix}_{stamp}.xlsx',
            content_type=(
                'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
            ),
        ), total_count, part_count

    output = tempfile.SpooledTemporaryFile(max_size=16 * 1024 * 1024, mode='w+b')
    with zipfile.ZipFile(output, mode='w', compression=zipfile.ZIP_DEFLATED) as archive:
        for part_number in range(1, part_count + 1):
            memories = list(islice(iterator, rows_per_file))
            workbook = _build_workbook(
                memories,
                total_count=total_count,
                part_number=part_number,
                part_count=part_count,
                scope_label=scope_label,
            )
            part_name = (
                f'{filename_prefix}_{stamp}_part_{part_number:03d}_of_{part_count:03d}.xlsx'
            )
            with tempfile.NamedTemporaryFile(suffix='.xlsx') as part_file:
                workbook.save(part_file.name)
                archive.write(part_file.name, arcname=part_name)
    output.seek(0)
    return FileResponse(
        output,
        as_attachment=True,
        filename=f'{filename_prefix}_{stamp}.zip',
        content_type='application/zip',
    ), total_count, part_count
