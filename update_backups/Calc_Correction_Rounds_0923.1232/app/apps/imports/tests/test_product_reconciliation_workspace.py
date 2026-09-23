from decimal import Decimal
from io import BytesIO
import zipfile

from django.contrib.auth import get_user_model
from django.test import TestCase
from django.urls import reverse
from openpyxl import Workbook, load_workbook

from apps.audit.models import AuditEvent
from apps.clients.models import Client
from apps.imports.models import (
    ExternalDataCorrectionMemory,
    ExternalDataFile,
    ProductReconciliationDecision,
    ProductReconciliationRule,
    ProductSourceRow,
)
from apps.imports.management.commands.import_sth_excel import Command as ImportSthCommand
from apps.imports.services.product_reconciliation_workspace import (
    ProductReconciliationWorkspaceError,
    build_workspace,
    derive_freight_type_from_pallet,
    preview_decision,
    save_product_reconciliation_decisions,
    save_reconciliation_rule,
)
from apps.imports.services.product_reconciliation_apply import (
    ProductReconciliationApplyBlocked,
    apply_product_reconciliation,
    build_product_apply_plan,
    create_recommended_product_drafts,
    rollback_latest_product_apply,
)
from apps.imports.services.product_reconciliation_memory import (
    ProductReconciliationMemoryError,
    adopt_operational_product_memories,
    backfill_applied_product_memories,
    prepare_exact_memory_drafts,
    reuse_all_exact_product_memories,
    reuse_product_memories,
)
from apps.imports.services.correction_memory_export import (
    correction_memory_export_response,
)
from apps.products.models import Product
from apps.saved_estimates.models import SavedEstimate


class ProductReconciliationWorkspaceTests(TestCase):
    def setUp(self):
        self.user = get_user_model().objects.create_user(
            username='reconciliation-admin',
            password='password',
            is_staff=True,
            is_superuser=True,
        )
        self.client_obj = Client.objects.create(
            code='STH',
            name='Stenhoj Australia',
            active=True,
        )
        self.external_file = ExternalDataFile.objects.create(
            client=self.client_obj,
            file_type='PRODUCTS',
            source_method='ADMIN_UPLOAD',
            original_filename='products.xls',
            status='VALIDATED',
            uploaded_by=self.user,
        )
        self.product = Product.objects.create(
            client=self.client_obj,
            sku='20772',
            name='120',
            description='80',
            length_m=Decimal('1.2000'),
            width_m=Decimal('0.8000'),
            height_m=Decimal('0.3500'),
            weight_kg=Decimal('70.0000'),
            cubic_m3=Decimal('0.336000'),
            freight_type='P',
        )
        ProductSourceRow.objects.create(
            external_file=self.external_file,
            source_row_number=6792,
            product_code_raw='20772',
            product_code_normalized='20772',
            name='Wheel Sup OffRd 500mm Max875kg',
            description='Wheel Support Device Offroad',
            category='TEST',
            length_mm=0,
            width_mm=0,
            height_mm=0,
            weight_kg=Decimal('45.2'),
            cubic_m3=Decimal('0.184'),
            quantity=1,
            pallet=1,
            source_status='L',
        )

    def product_snapshot(self):
        return list(Product.objects.order_by('pk').values())

    def clone_source_file(self, *, filename='products-copy.xls', source_name=None, sha256=''):
        external_file = ExternalDataFile.objects.create(
            client=self.client_obj,
            file_type='PRODUCTS',
            source_method='ADMIN_UPLOAD',
            original_filename=filename,
            status='VALIDATED',
            uploaded_by=self.user,
            sha256=sha256,
        )
        original = ProductSourceRow.objects.get(external_file=self.external_file)
        ProductSourceRow.objects.create(
            external_file=external_file,
            source_row_number=original.source_row_number,
            product_code_raw=original.product_code_raw,
            product_code_normalized=original.product_code_normalized,
            name=source_name if source_name is not None else original.name,
            description=original.description,
            category=original.category,
            length_mm=original.length_mm,
            width_mm=original.width_mm,
            height_mm=original.height_mm,
            weight_kg=original.weight_kg,
            cubic_m3=original.cubic_m3,
            quantity=original.quantity,
            pallet=original.pallet,
            comment=original.comment,
            source_status=original.source_status,
        )
        return external_file

    def test_workspace_groups_source_zero_dimensions(self):
        rows, summary = build_workspace(self.external_file)

        self.assertEqual(rows[0]['group_key'], 'SOURCE_DIMENSIONS_ZERO')
        self.assertEqual(summary['group_counts']['SOURCE_DIMENSIONS_ZERO'], 1)
        self.assertEqual(summary['draft_decisions'], 0)
        self.assertEqual(summary['memory_new_issue'], 1)

    def test_workspace_counts_active_reusable_rule_matches(self):
        save_reconciliation_rule(
            self.external_file,
            name='Protect valid Calculator dimensions',
            group_key='SOURCE_DIMENSIONS_ZERO',
            field_decisions={
                'name': 'SOURCE',
                'description': 'SOURCE',
                'dimensions': 'OPERATIONAL',
                'weight': 'OPERATIONAL',
                'cubic': 'OPERATIONAL',
                'freight_type': 'SOURCE',
            },
            actor=self.user,
        )

        rows, summary = build_workspace(self.external_file)

        self.assertEqual(summary['memory_rule_match'], 1)
        self.assertEqual(rows[0]['matching_rules'][0].name, 'Protect valid Calculator dimensions')

    def test_case_pallet_rule_uses_zero_and_any_positive_number(self):
        self.assertEqual(derive_freight_type_from_pallet(0), ('C', None))
        self.assertEqual(derive_freight_type_from_pallet(1), ('P', None))
        self.assertEqual(derive_freight_type_from_pallet(7), ('P', None))
        self.assertEqual(
            derive_freight_type_from_pallet(None),
            (None, 'PALLET_VALUE_MISSING'),
        )
        self.assertEqual(
            derive_freight_type_from_pallet(-1),
            (None, 'PALLET_VALUE_NEGATIVE'),
        )
        self.assertEqual(
            derive_freight_type_from_pallet('invalid'),
            (None, 'PALLET_VALUE_INVALID'),
        )

    def test_invalid_pallet_requires_manual_freight_type(self):
        source_row = ProductSourceRow.objects.get(external_file=self.external_file)
        source_row.pallet = None
        source_row.save(update_fields=['pallet'])

        rows, summary = build_workspace(self.external_file)

        self.assertTrue(rows[0]['freight_type_review_required'])
        self.assertEqual(rows[0]['group_key'], 'FREIGHT_TYPE_REVIEW')
        self.assertEqual(summary['freight_type_review_required'], 1)
        with self.assertRaisesMessage(
            ProductReconciliationWorkspaceError,
            'Select C or P manually',
        ):
            save_product_reconciliation_decisions(
                self.external_file,
                skus=['20772'],
                field_decisions={
                    'name': 'OPERATIONAL',
                    'description': 'OPERATIONAL',
                    'dimensions': 'OPERATIONAL',
                    'weight': 'OPERATIONAL',
                    'cubic': 'OPERATIONAL',
                    'freight_type': 'SOURCE',
                },
                actor=self.user,
            )

        decision = save_product_reconciliation_decisions(
            self.external_file,
            skus=['20772'],
            field_decisions={
                'name': 'OPERATIONAL',
                'description': 'OPERATIONAL',
                'dimensions': 'OPERATIONAL',
                'weight': 'OPERATIONAL',
                'cubic': 'OPERATIONAL',
                'freight_type': 'CUSTOM',
            },
            custom_values={'freight_type': 'C'},
            actor=self.user,
        )[0]
        self.assertEqual(preview_decision(rows[0], decision)['proposed_values']['freight_type'], 'C')

    def test_bulk_draft_and_preview_do_not_change_product_or_freight_type(self):
        before = self.product_snapshot()
        decisions = save_product_reconciliation_decisions(
            self.external_file,
            skus=['20772'],
            field_decisions={
                'name': 'SOURCE',
                'description': 'SOURCE',
                'dimensions': 'OPERATIONAL',
                'weight': 'OPERATIONAL',
                'cubic': 'OPERATIONAL',
            },
            actor=self.user,
        )
        rows, summary = build_workspace(self.external_file)
        preview = preview_decision(rows[0], decisions[0])

        self.assertEqual(summary['draft_decisions'], 1)
        self.assertEqual(preview['proposed_values']['name'], 'Wheel Sup OffRd 500mm Max875kg')
        self.assertEqual(preview['proposed_values']['length_m'], Decimal('1.2000'))
        self.assertEqual(preview['proposed_values']['freight_type'], 'P')
        self.assertTrue(preview['operational_update_available'])
        self.assertEqual(self.product_snapshot(), before)
        self.assertTrue(AuditEvent.objects.filter(
            event_type='PRODUCT_RECONCILIATION_DRAFT_SAVED'
        ).exists())

    def test_individual_custom_values_remain_draft_only(self):
        before = self.product_snapshot()
        decision = save_product_reconciliation_decisions(
            self.external_file,
            skus=['20772'],
            field_decisions={
                'name': 'CUSTOM',
                'description': 'OPERATIONAL',
                'dimensions': 'CUSTOM',
                'weight': 'OPERATIONAL',
                'cubic': 'CUSTOM',
            },
            custom_values={
                'name': 'Reviewed wheel support',
                'length_m': '1.25',
                'width_m': '0.80',
                'height_m': '0.35',
                'cubic_m3': '0.350',
            },
            actor=self.user,
        )[0]

        self.assertEqual(decision.custom_values['name'], 'Reviewed wheel support')
        self.assertEqual(decision.custom_values['length_m'], '1.25')
        self.assertEqual(self.product_snapshot(), before)

    def test_reusable_rule_is_saved_without_applying_operational_changes(self):
        before = self.product_snapshot()
        rule = save_reconciliation_rule(
            self.external_file,
            name='Keep Calculator when source dimensions are zero',
            group_key='SOURCE_DIMENSIONS_ZERO',
            field_decisions={
                'name': 'SOURCE',
                'description': 'SOURCE',
                'dimensions': 'OPERATIONAL',
                'weight': 'OPERATIONAL',
                'cubic': 'OPERATIONAL',
            },
            actor=self.user,
        )

        self.assertTrue(rule.active)
        self.assertEqual(ProductReconciliationRule.objects.count(), 1)
        self.assertEqual(self.product_snapshot(), before)

    def test_workspace_screen_saves_selected_bulk_draft(self):
        before = self.product_snapshot()
        self.client.force_login(self.user)
        url = reverse(
            'admin:imports_externaldatafile_product_reconciliation',
            args=[self.external_file.pk],
        )

        response = self.client.post(url, {
            'workspace_action': 'save_bulk',
            'group': 'SOURCE_DIMENSIONS_ZERO',
            'selected_skus': ['20772'],
            'preset': 'KEEP_PHYSICAL_USE_SOURCE_TEXT',
            'scope': 'selected',
            'notes': 'Reviewed as a group.',
        })

        self.assertEqual(response.status_code, 302)
        decision = ProductReconciliationDecision.objects.get(
            external_file=self.external_file,
            product_code_normalized='20772',
        )
        self.assertEqual(decision.field_decisions['dimensions'], 'OPERATIONAL')
        self.assertEqual(self.product_snapshot(), before)

        response = self.client.get(f'{url}?group=SOURCE_DIMENSIONS_ZERO&mode=preview')
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'Preview draft decisions')
        self.assertContains(response, 'C/P proposal')
        self.assertContains(response, '0 → C')
        self.assertContains(response, 'Apply 1 reviewed change(s)')

    def test_recommended_correction_applies_source_text_and_protected_cp_then_rolls_back(self):
        old_name = self.product.name
        old_description = self.product.description
        old_weight = self.product.weight_kg

        decisions = create_recommended_product_drafts(
            self.external_file,
            actor=self.user,
        )
        self.assertEqual(len(decisions), 1)
        plan = build_product_apply_plan(self.external_file.pk)
        self.assertTrue(plan['can_apply'])
        self.assertEqual(plan['update_count'], 1)

        batch = apply_product_reconciliation(self.external_file.pk, actor=self.user)
        self.product.refresh_from_db()
        self.assertEqual(self.product.name, 'Wheel Sup OffRd 500mm Max875kg')
        self.assertEqual(self.product.description, 'Wheel Support Device Offroad')
        self.assertEqual(self.product.weight_kg, old_weight)
        self.assertEqual(self.product.freight_type, 'P')
        self.assertTrue(batch['batch_id'])
        self.assertTrue(AuditEvent.objects.filter(
            event_type='PRODUCT_RECONCILIATION_APPLIED'
        ).exists())
        memory = ExternalDataCorrectionMemory.objects.get(record_key='20772')
        self.assertTrue(memory.is_active)
        self.assertEqual(memory.approved_data['approved_values']['name'], self.product.name)

        rollback_latest_product_apply(self.external_file.pk, actor=self.user)
        self.product.refresh_from_db()
        self.assertEqual(self.product.name, old_name)
        self.assertEqual(self.product.description, old_description)
        self.assertTrue(AuditEvent.objects.filter(
            event_type='PRODUCT_RECONCILIATION_ROLLED_BACK'
        ).exists())
        memory.refresh_from_db()
        self.assertFalse(memory.is_active)

    def test_applied_workspace_shows_statuses_and_blocks_new_drafts(self):
        create_recommended_product_drafts(self.external_file, actor=self.user)
        apply_product_reconciliation(self.external_file.pk, actor=self.user)
        rows, summary = build_workspace(self.external_file)
        self.assertEqual(len(rows), 1)
        self.assertEqual(summary['draft_decisions'], 0)
        self.assertEqual(summary['ready_decisions'], 0)
        self.assertEqual(summary['applied_decisions'], 1)
        self.assertEqual(summary['reusable_previous_solutions'], 0)

        self.client.force_login(self.user)
        url = reverse(
            'admin:imports_externaldatafile_product_reconciliation',
            args=[self.external_file.pk],
        )
        response = self.client.get(f'{url}?mode=preview')
        self.assertContains(response, 'No pending drafts. 1 decision was already applied.')
        self.assertContains(response, 'Rollback latest Apply')
        self.assertContains(response, 'Download memory Excel')
        self.assertNotContains(response, 'Prepare recommended initial correction')
        self.assertNotContains(response, 'Reuse all previous approved solutions')
        self.assertNotContains(response, 'Apply 1 reviewed change(s)')

        for action in ('create_recommended', 'reuse_all_exact_memory', 'save_bulk'):
            response = self.client.post(url, {'workspace_action': action, 'group': 'ALL_DIFFERENCES'})
            self.assertEqual(response.status_code, 302)
        with self.assertRaises(ProductReconciliationApplyBlocked):
            create_recommended_product_drafts(self.external_file, actor=self.user)
        with self.assertRaises(ProductReconciliationWorkspaceError):
            save_product_reconciliation_decisions(
                self.external_file,
                skus=['20772'],
                field_decisions={'name': 'SOURCE'},
                actor=self.user,
            )
        decision = ProductReconciliationDecision.objects.get(external_file=self.external_file)
        self.assertEqual(decision.decision_status, 'APPLIED')
        self.assertEqual(build_product_apply_plan(self.external_file.pk)['decision_count'], 0)

        rollback_latest_product_apply(self.external_file.pk, actor=self.user)
        rows, summary = build_workspace(self.external_file)
        self.assertEqual(summary['ready_decisions'], 1)
        self.assertEqual(summary['applied_decisions'], 0)
        response = self.client.get(f'{url}?mode=preview')
        self.assertContains(response, 'Apply 1 reviewed change(s)')

    def test_source_only_rows_do_not_count_as_new_operational_issues(self):
        ProductSourceRow.objects.create(
            external_file=self.external_file,
            source_row_number=8000,
            product_code_raw='SOURCE-ONLY',
            product_code_normalized='SOURCE-ONLY',
            name='Source reference',
            description='',
        )
        rows, summary = build_workspace(self.external_file)
        self.assertEqual(summary['source_only'], 1)
        self.assertEqual(summary['memory_new_issue'], 1)
        self.assertEqual(len([row for row in rows if row['status'] == 'SOURCE_ONLY']), 1)

    def test_exact_previous_problem_prepares_draft_without_operational_update(self):
        create_recommended_product_drafts(self.external_file, actor=self.user)
        apply_product_reconciliation(self.external_file.pk, actor=self.user)
        repeated = self.clone_source_file()
        before = self.product_snapshot()

        rows, summary = build_workspace(repeated)
        self.assertEqual(rows[0]['memory_state'], 'EXACT')
        self.assertEqual(summary['memory_exact'], 1)
        self.assertEqual(summary['reusable_previous_solutions'], 1)
        prepared = prepare_exact_memory_drafts(repeated, actor=self.user)

        self.assertEqual(len(prepared), 1)
        self.assertEqual(prepared[0].decision_status, 'DRAFT')
        self.assertIn('Reused approved reconciliation memory', prepared[0].notes)
        self.assertEqual(self.product_snapshot(), before)

    def test_reuse_all_exact_previous_solutions_refreshes_drafts_only(self):
        create_recommended_product_drafts(self.external_file, actor=self.user)
        apply_product_reconciliation(self.external_file.pk, actor=self.user)
        repeated = self.clone_source_file()
        before = self.product_snapshot()

        prepared = reuse_all_exact_product_memories(repeated, actor=self.user)

        self.assertEqual(len(prepared), 1)
        self.assertEqual(prepared[0].decision_status, 'DRAFT')
        self.assertEqual(self.product_snapshot(), before)
        self.assertTrue(AuditEvent.objects.filter(
            event_type='PRODUCT_RECONCILIATION_ALL_EXACT_MEMORIES_REUSED'
        ).exists())

    def test_changed_source_warns_and_requires_explicit_memory_reuse(self):
        create_recommended_product_drafts(self.external_file, actor=self.user)
        apply_product_reconciliation(self.external_file.pk, actor=self.user)
        changed = self.clone_source_file(source_name='Updated source description')

        rows, summary = build_workspace(changed)
        self.assertEqual(rows[0]['memory_state'], 'SOURCE_CHANGED')
        self.assertEqual(summary['memory_source_changed'], 1)
        self.assertEqual(prepare_exact_memory_drafts(changed, actor=self.user), [])
        self.assertFalse(ProductReconciliationDecision.objects.filter(
            external_file=changed
        ).exists())

        reused = reuse_product_memories(
            changed,
            skus=['20772'],
            actor=self.user,
        )
        self.assertEqual(len(reused), 1)
        self.assertEqual(reused[0].decision_status, 'DRAFT')

    def test_adopt_current_calculator_requires_reason_and_does_not_update_product(self):
        before = self.product_snapshot()
        with self.assertRaises(ProductReconciliationMemoryError):
            adopt_operational_product_memories(
                self.external_file,
                skus=['20772'],
                approval_note='',
                actor=self.user,
            )

        memories = adopt_operational_product_memories(
            self.external_file,
            skus=['20772'],
            approval_note='Validated against the initial Calculator catalogue.',
            actor=self.user,
        )
        self.assertEqual(len(memories), 1)
        self.assertTrue(memories[0].approved_data['adopted_operational_state'])
        self.assertEqual(self.product_snapshot(), before)

    def test_pre_feature_applied_decision_is_backfilled_as_memory(self):
        create_recommended_product_drafts(self.external_file, actor=self.user)
        apply_product_reconciliation(self.external_file.pk, actor=self.user)
        ExternalDataCorrectionMemory.objects.all().delete()

        memories = backfill_applied_product_memories(
            self.external_file,
            actor=self.user,
        )

        self.assertEqual(len(memories), 1)
        self.assertTrue(memories[0].approved_data['backfilled_from_applied_decision'])
        self.assertTrue(AuditEvent.objects.filter(
            event_type='PRODUCT_RECONCILIATION_MEMORY_BACKFILLED'
        ).exists())

    def test_memory_download_button_exports_reviewable_excel(self):
        adopt_operational_product_memories(
            self.external_file,
            skus=['20772'],
            approval_note='Approved historical Calculator value.',
            actor=self.user,
        )
        self.client.force_login(self.user)
        url = reverse(
            'admin:imports_externaldatafile_download_product_memory',
            args=[self.external_file.pk],
        )

        response = self.client.get(url)
        content = b''.join(response.streaming_content)
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
        rows = list(workbook['Memories'].iter_rows(values_only=True))

        self.assertEqual(response.status_code, 200)
        self.assertIn('.xlsx', response['Content-Disposition'])
        self.assertEqual(rows[0][3], 'SKU / record key')
        self.assertEqual(rows[1][3], '20772')
        self.assertTrue(AuditEvent.objects.filter(
            event_type='PRODUCT_RECONCILIATION_MEMORY_EXPORTED'
        ).exists())

        admin_response = self.client.get(reverse(
            'admin:imports_externaldatacorrectionmemory_download_excel'
        ))
        admin_content = b''.join(admin_response.streaming_content)
        admin_workbook = load_workbook(
            BytesIO(admin_content), read_only=True, data_only=True
        )
        self.assertEqual(admin_response.status_code, 200)
        self.assertEqual(admin_workbook['Summary']['B2'].value, 'Reconciliation Decision Memory')

    def test_large_memory_export_is_split_into_numbered_excel_files(self):
        adopt_operational_product_memories(
            self.external_file,
            skus=['20772'],
            approval_note='First approved value.',
            actor=self.user,
        )
        second_product = Product.objects.create(
            client=self.client_obj,
            sku='SECOND',
            name='Second product',
            description='',
            length_m=Decimal('1'),
            width_m=Decimal('1'),
            height_m=Decimal('1'),
            weight_kg=Decimal('1'),
            cubic_m3=Decimal('1'),
            freight_type='C',
        )
        self.assertTrue(second_product.pk)
        adopt_operational_product_memories(
            self.external_file,
            skus=['20772'],
            approval_note='Updated first value.',
            actor=self.user,
        )
        second_file = ExternalDataFile.objects.create(
            client=self.client_obj,
            file_type='PRODUCTS',
            source_method='ADMIN_UPLOAD',
            original_filename='empty-products.xls',
            status='VALIDATED',
            uploaded_by=self.user,
        )
        adopt_operational_product_memories(
            second_file,
            skus=['SECOND'],
            approval_note='Second approved value.',
            actor=self.user,
        )

        response, record_count, part_count = correction_memory_export_response(
            ExternalDataCorrectionMemory.objects.filter(
                workflow_key='PRODUCT_RECONCILIATION'
            ),
            rows_per_file=1,
        )
        content = b''.join(response.streaming_content)
        with zipfile.ZipFile(BytesIO(content)) as archive:
            names = archive.namelist()

        self.assertEqual(record_count, 2)
        self.assertEqual(part_count, 2)
        self.assertEqual(len(names), 2)
        self.assertTrue(all(name.endswith('.xlsx') for name in names))

    def test_operational_only_removal_is_blocked_when_saved_quote_references_sku(self):
        ProductSourceRow.objects.filter(external_file=self.external_file).delete()
        SavedEstimate.objects.create(
            reference='FQ-STH-TEST-1',
            client=self.client_obj,
            created_by=self.user,
            created_by_label='tester',
            input_snapshot={'lines': [{'sku': '20772'}]},
            result_snapshot=[],
        )
        create_recommended_product_drafts(self.external_file, actor=self.user)
        plan = build_product_apply_plan(self.external_file.pk)

        self.assertFalse(plan['can_apply'])
        self.assertEqual(plan['delete_count'], 1)
        self.assertEqual(plan['items'][0]['reference_summary']['saved_estimates'], 1)
        with self.assertRaises(ProductReconciliationApplyBlocked):
            apply_product_reconciliation(self.external_file.pk, actor=self.user)
        self.assertTrue(Product.objects.filter(pk=self.product.pk).exists())

    def test_operational_only_unreferenced_product_is_removed_and_restored(self):
        ProductSourceRow.objects.filter(external_file=self.external_file).delete()
        create_recommended_product_drafts(self.external_file, actor=self.user)

        plan = build_product_apply_plan(self.external_file.pk)
        self.assertTrue(plan['can_apply'])
        self.assertEqual(plan['delete_count'], 1)
        apply_product_reconciliation(self.external_file.pk, actor=self.user)
        self.assertFalse(Product.objects.filter(pk=self.product.pk).exists())

        rollback_latest_product_apply(self.external_file.pk, actor=self.user)
        restored = Product.objects.get(pk=self.product.pk)
        self.assertEqual(restored.sku, '20772')
        self.assertEqual(restored.name, '120')

    def test_dimension_scale_warning_blocks_source_dimension_apply(self):
        source = ProductSourceRow.objects.get(external_file=self.external_file)
        source.length_mm = Decimal('120')
        source.width_mm = Decimal('80')
        source.height_mm = Decimal('35')
        source.save(update_fields=['length_mm', 'width_mm', 'height_mm'])
        rows, summary = build_workspace(self.external_file)

        self.assertTrue(rows[0]['dimension_unit_review_required'])
        self.assertEqual(summary['group_counts']['DIMENSION_UNIT_REVIEW'], 1)
        save_product_reconciliation_decisions(
            self.external_file,
            skus=['20772'],
            field_decisions={
                'name': 'SOURCE',
                'description': 'SOURCE',
                'dimensions': 'SOURCE',
                'weight': 'OPERATIONAL',
                'cubic': 'OPERATIONAL',
                'freight_type': 'SOURCE',
            },
            actor=self.user,
        )
        plan = build_product_apply_plan(self.external_file.pk)
        self.assertFalse(plan['can_apply'])
        self.assertIn('manual unit review', plan['blockers'][0])

    def test_cp_difference_count_is_transversal_to_physical_group(self):
        self.product.freight_type = 'C'
        self.product.save(update_fields=['freight_type'])
        rows, summary = build_workspace(self.external_file)

        self.assertEqual(rows[0]['group_key'], 'SOURCE_DIMENSIONS_ZERO')
        self.assertEqual(summary['group_counts']['FREIGHT_TYPE_DIFFERENCES'], 1)

    def test_legacy_workbook_bootstrap_never_maps_dimensions_to_product_text(self):
        Product.objects.all().delete()
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = 'SKUs'
        worksheet.append([
            'SKU', 'Length (cm)', 'Width (cm)', 'Height (cm)',
            'Length (m)', 'Width (m)', 'Height (m)', 'Weight', 'Cubic', 'Type',
        ])
        worksheet.append([
            'SKU-BOOT', 120, 80, 35, 1.2, 0.8, 0.35, 70, 0.336, 'P',
        ])

        ImportSthCommand().import_products(workbook, self.client_obj)
        product = Product.objects.get(sku='SKU-BOOT')
        self.assertEqual(product.name, 'SKU-BOOT')
        self.assertEqual(product.description, '')

    def test_edit_action_is_visible_beside_sku_and_targets_editor(self):
        self.client.force_login(self.user)
        url = reverse(
            'admin:imports_externaldatafile_product_reconciliation',
            args=[self.external_file.pk],
        )

        response = self.client.get(f'{url}?group=SOURCE_DIMENSIONS_ZERO')
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'Edit product')
        self.assertContains(response, '#individual-editor')
        self.assertNotContains(response, 'Reuse all previous approved solutions')
        self.assertContains(response, 'Download memory Excel')

        response = self.client.get(
            f'{url}?group=SOURCE_DIMENSIONS_ZERO&edit=20772'
        )
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'id="individual-editor"', html=False)
        self.assertContains(response, 'Individual review · 20772')
